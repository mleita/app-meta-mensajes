#!/usr/bin/env python3
"""
% captado por habitualistas — transferencias de usados, provincia de Córdoba.

Fuente única: DNRPA / Ministerio de Justicia, dataset "Transferencias de autos"
  https://datos.jus.gob.ar/dataset/transferencias-de-autos
  (API CKAN: https://datos.jus.gob.ar/api/3/action/package_show?id=transferencias-de-autos)

Indicador por mes:
  transferencias_totales        = filas con registro_seccional_provincia == CORDOBA
  transferencias_habitualistas  = ídem y tramite_tipo == "TRANSFERENCIA A COMERCIANTE HABITUALISTA"
  porcentaje_habitualistas      = habitualistas / totales * 100

Uso:
  python dnrpa_habitualistas_cordoba.py                 # descarga (API CKAN) + procesa
  python dnrpa_habitualistas_cordoba.py --solo-local    # procesa lo que haya en data/raw
  python dnrpa_habitualistas_cordoba.py --desde 2025-01 --hasta 2026-12

Para actualizar cada mes basta volver a correrlo: descarga sólo los archivos nuevos.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

import pandas as pd

BASE = Path(__file__).resolve().parent
RAW = BASE / "data" / "raw"
OUT = BASE / "output"

CKAN_API = "https://datos.jus.gob.ar/api/3/action/package_show?id=transferencias-de-autos"
TIPO_HABITUALISTA = "TRANSFERENCIA A COMERCIANTE HABITUALISTA"
PROVINCIA_OBJETIVO = "CORDOBA"  # comparado tras normalizar (mayúsculas, sin tildes, sin espacios extra)

# Sanity check externo aportado por el usuario (NO es fuente; sólo comparación).
REFERENCIA_EXTERNA_2025 = {
    1: 19608, 2: 17982, 3: 16955, 4: 19142, 5: 18678, 6: 17931,
    7: 21573, 8: 20549, 9: 20957, 10: 21059, 11: 16497, 12: 18009,
}

MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
PATRON_ARCHIVO = re.compile(r"dnrpa-transferencias-autos-(\d{4})(\d{2})?\.(csv|zip)$", re.I)
PALABRAS_SOSPECHOSAS = ["ANUL", "RECTIF", "BAJA", "DESIST", "REVOC", "CANCEL", "OBSERV", "RECHAZ"]


def norm(s) -> str:
    if pd.isna(s):
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


# --------------------------------------------------------------------------- descarga
def listar_recursos(raw: Path):
    import requests

    r = requests.get(CKAN_API, timeout=60)
    r.raise_for_status()
    pkg = r.json()["result"]
    recursos = []
    for res in pkg["resources"]:
        url = res.get("url") or ""
        m = PATRON_ARCHIVO.search(url.split("?")[0])
        if not m:
            continue
        recursos.append({
            "nombre": res.get("name"),
            "url": url,
            "archivo": Path(url.split("?")[0]).name,
            "anio": int(m.group(1)),
            "mes": int(m.group(2)) if m.group(2) else None,
            "last_modified": res.get("last_modified") or res.get("created"),
        })
    meta = {k: pkg.get(k) for k in ("title", "metadata_modified", "notes")}
    (raw / "ckan_package_show.json").write_text(json.dumps(pkg, ensure_ascii=False, indent=2), "utf-8")
    return recursos, meta


def descargar(recursos: list[dict], desde: tuple, hasta: tuple, raw: Path) -> list[dict]:
    import requests

    elegidos = []
    for r in recursos:
        if r["mes"] is not None:
            if not (desde <= (r["anio"], r["mes"]) <= hasta):
                continue
        elif not (desde[0] <= r["anio"] <= hasta[0]):
            continue
        destino = raw / r["archivo"]
        if not destino.exists():
            print(f"  descargando {r['url']}")
            with requests.get(r["url"], stream=True, timeout=600) as resp:
                resp.raise_for_status()
                tmp = destino.with_suffix(destino.suffix + ".part")
                with open(tmp, "wb") as f:
                    for chunk in resp.iter_content(1 << 20):
                        f.write(chunk)
                tmp.rename(destino)
        else:
            print(f"  ya existe {destino.name}")
        elegidos.append(r)
    return elegidos


def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# --------------------------------------------------------------------------- lectura
def abrir_bytes(p: Path) -> list[tuple[str, bytes]]:
    """Devuelve [(nombre_csv, contenido)] — soporta CSV sueltos o ZIP con CSVs."""
    if p.suffix.lower() == ".zip":
        with zipfile.ZipFile(p) as z:
            return [(n, z.read(n)) for n in z.namelist() if n.lower().endswith(".csv")]
    return [(p.name, p.read_bytes())]


def detectar_formato(contenido: bytes) -> dict:
    encoding = "utf-8-sig" if contenido.startswith(b"\xef\xbb\xbf") else "utf-8"
    try:
        muestra = contenido[:200_000].decode(encoding)
    except UnicodeDecodeError:
        encoding = "latin-1"
        muestra = contenido[:200_000].decode(encoding)
    try:
        sep = csv.Sniffer().sniff(muestra.split("\n", 1)[0], delimiters=",;|\t").delimiter
    except csv.Error:
        sep = ","
    return {"encoding": encoding, "separador": sep}


def formato_fecha(serie: pd.Series) -> str:
    ej = serie.dropna().astype(str).head(1000)
    if ej.empty:
        return "sin datos"
    if ej.str.fullmatch(r"\d{4}-\d{2}-\d{2}").all():
        return "AAAA-MM-DD"
    if ej.str.fullmatch(r"\d{2}/\d{2}/\d{4}").all():
        return "DD/MM/AAAA"
    return "mixto: " + ", ".join(ej.head(3))


def parsear_fecha(serie: pd.Series) -> pd.Series:
    f = pd.to_datetime(serie, format="%Y-%m-%d", errors="coerce")
    faltan = f.isna() & serie.notna()
    if faltan.any():
        f[faltan] = pd.to_datetime(serie[faltan], dayfirst=True, errors="coerce")
    return f


def mes_de_archivo(nombre: str):
    m = PATRON_ARCHIVO.search(nombre)
    return (int(m.group(1)), int(m.group(2))) if m and m.group(2) else None


def procesar_archivos(archivos: list[Path]):
    inspeccion, cordoba_partes = [], []
    tipos_pais = {}
    provincias_crudas = {}
    for p in archivos:
        for nombre, contenido in abrir_bytes(p):
            fmt = detectar_formato(contenido)
            df = pd.read_csv(io.BytesIO(contenido), sep=fmt["separador"], encoding=fmt["encoding"],
                             dtype=str, keep_default_na=True)
            df.columns = [c.strip() for c in df.columns]
            faltan = [c for c in ("tramite_fecha", "tramite_tipo", "registro_seccional_provincia")
                      if c not in df.columns]
            if faltan:
                raise SystemExit(f"{nombre}: faltan columnas {faltan}. Columnas: {list(df.columns)}")

            for v, n in df["registro_seccional_provincia"].value_counts(dropna=False).items():
                provincias_crudas[v] = provincias_crudas.get(v, 0) + int(n)
            for v, n in df["tramite_tipo"].value_counts(dropna=False).items():
                tipos_pais[v] = tipos_pais.get(v, 0) + int(n)

            es_cba = df["registro_seccional_provincia"].map(norm) == PROVINCIA_OBJETIVO
            cba = df[es_cba].copy()
            cba["archivo_origen"] = nombre
            fam = mes_de_archivo(nombre)
            cba["archivo_anio"], cba["archivo_mes"] = (fam if fam else (None, None))
            cordoba_partes.append(cba)

            fechas = parsear_fecha(df["tramite_fecha"])
            inspeccion.append({
                "archivo": nombre,
                "bytes": len(contenido),
                "encoding": fmt["encoding"],
                "separador": fmt["separador"],
                "filas_pais": len(df),
                "filas_cordoba": int(es_cba.sum()),
                "columnas": list(df.columns),
                "formato_tramite_fecha": formato_fecha(df["tramite_fecha"]),
                "tramite_fecha_min": str(fechas.min().date()) if fechas.notna().any() else None,
                "tramite_fecha_max": str(fechas.max().date()) if fechas.notna().any() else None,
                "tramite_fecha_no_parseable": int((fechas.isna() & df["tramite_fecha"].notna()).sum()),
            })
            print(f"  {nombre}: {len(df):,} filas país, {int(es_cba.sum()):,} Córdoba")
    cordoba = pd.concat(cordoba_partes, ignore_index=True) if cordoba_partes else pd.DataFrame()
    return cordoba, inspeccion, tipos_pais, provincias_crudas


# --------------------------------------------------------------------------- cálculo
def calcular(cba: pd.DataFrame, desde: tuple, hasta: tuple, eliminar_duplicados: bool):
    control = {}
    cba = cba.copy()
    cba["fecha"] = parsear_fecha(cba["tramite_fecha"])
    cba["anio"] = cba["fecha"].dt.year.astype("Int64")
    cba["mes"] = cba["fecha"].dt.month.astype("Int64")

    cols_datos = [c for c in cba.columns if c not in ("archivo_origen", "archivo_anio", "archivo_mes",
                                                       "fecha", "anio", "mes")]
    dup_mask = cba.duplicated(subset=cols_datos, keep="first")
    control["filas_cordoba_leidas"] = len(cba)
    control["filas_duplicadas_exactas"] = int(dup_mask.sum())
    # mismo contenido presente en dos archivos distintos (p.ej. mensual + anual)
    dup_entre_archivos = cba[cba.duplicated(subset=cols_datos, keep=False)] \
        .groupby(cols_datos, dropna=False)["archivo_origen"].nunique()
    control["duplicados_entre_archivos"] = int((dup_entre_archivos > 1).sum())
    if eliminar_duplicados:
        cba = cba[~dup_mask]
    control["duplicados_eliminados"] = bool(eliminar_duplicados)

    control["tramite_tipo_nulos"] = int(cba["tramite_tipo"].isna().sum() +
                                        (cba["tramite_tipo"].fillna("x").str.strip() == "").sum())
    control["tramite_fecha_nulas_o_invalidas"] = int(cba["fecha"].isna().sum())

    tipos = cba["tramite_tipo"].fillna("<NULO>").value_counts()
    control["tramite_tipo_cordoba"] = {k: int(v) for k, v in tipos.items()}
    control["tipos_con_HABITUALISTA"] = {k: int(v) for k, v in tipos.items() if "HABITUALISTA" in norm(k)}
    control["tipos_sospechosos_anulacion_rectificacion"] = {
        k: int(v) for k, v in tipos.items() if any(w in norm(k) for w in PALABRAS_SOSPECHOSAS)}

    # tramite_fecha vs mes del archivo (el archivo AAAAMM = trámites inscriptos ese mes)
    con_mes = cba.dropna(subset=["archivo_mes"])
    if not con_mes.empty:
        distinto = (con_mes["anio"] != con_mes["archivo_anio"]) | (con_mes["mes"] != con_mes["archivo_mes"])
        control["filas_con_tramite_fecha_fuera_del_mes_del_archivo"] = int(distinto.sum())

    periodo = cba[(cba["anio"] * 100 + cba["mes"] >= desde[0] * 100 + desde[1]) &
                  (cba["anio"] * 100 + cba["mes"] <= hasta[0] * 100 + hasta[1])].copy()
    control["filas_fuera_de_periodo_descartadas"] = int(len(cba) - len(periodo))
    periodo["es_habitualista"] = periodo["tramite_tipo"].map(lambda x: str(x).strip()) == TIPO_HABITUALISTA

    # numerador ⊆ denominador: por construcción ambos salen del mismo DataFrame filtrado
    num = periodo[periodo["es_habitualista"]]
    control["numerador_subconjunto_denominador"] = bool(num.index.isin(periodo.index).all())

    mensual = (periodo.groupby(["anio", "mes"])
               .agg(transferencias_totales=("es_habitualista", "size"),
                    transferencias_habitualistas=("es_habitualista", "sum"))
               .reset_index())
    mensual["anio"] = mensual["anio"].astype(int)
    mensual["mes"] = mensual["mes"].astype(int)
    mensual["transferencias_habitualistas"] = mensual["transferencias_habitualistas"].astype(int)
    mensual["porcentaje_habitualistas"] = (mensual["transferencias_habitualistas"] /
                                           mensual["transferencias_totales"] * 100).round(2)
    return mensual, control, periodo


def resumen_anual(mensual: pd.DataFrame) -> pd.DataFrame:
    filas = []
    for anio, g in mensual.groupby("anio"):
        filas.append(_fila_resumen(f"Total {anio}" + ("" if len(g) == 12 else f" (acum. {len(g)} meses)"), g))
    filas.append(_fila_resumen("Período completo", mensual))
    return pd.DataFrame(filas)


def _fila_resumen(etiqueta, g):
    tot, hab = int(g["transferencias_totales"].sum()), int(g["transferencias_habitualistas"].sum())
    return {
        "periodo": etiqueta,
        "meses": len(g),
        "desde": f"{MESES[int(g.iloc[0]['mes']) - 1]} {int(g.iloc[0]['anio'])}",
        "hasta": f"{MESES[int(g.iloc[-1]['mes']) - 1]} {int(g.iloc[-1]['anio'])}",
        "transferencias_totales": tot,
        "transferencias_habitualistas": hab,
        "pct_acumulado": round(hab / tot * 100, 2) if tot else None,
        "promedio_simple_pct_mensual": round(g["porcentaje_habitualistas"].mean(), 2),
    }


def validacion_externa(mensual: pd.DataFrame) -> pd.DataFrame:
    m25 = mensual[mensual["anio"] == 2025].set_index("mes")["transferencias_totales"]
    filas = []
    for mes, ref in REFERENCIA_EXTERNA_2025.items():
        dn = int(m25.get(mes)) if mes in m25.index else None
        filas.append({"mes": f"{MESES[mes - 1]} 2025", "dnrpa_datos_abiertos": dn, "referencia_externa": ref,
                      "diferencia": (dn - ref) if dn is not None else None,
                      "diferencia_pct": round((dn - ref) / ref * 100, 2) if dn is not None else None})
    df = pd.DataFrame(filas)
    tot_dn = df["dnrpa_datos_abiertos"].sum(min_count=1)
    tot_ref = sum(REFERENCIA_EXTERNA_2025.values())
    df.loc[len(df)] = {"mes": "Total 2025", "dnrpa_datos_abiertos": tot_dn, "referencia_externa": tot_ref,
                       "diferencia": (tot_dn - tot_ref) if pd.notna(tot_dn) else None,
                       "diferencia_pct": round((tot_dn - tot_ref) / tot_ref * 100, 2) if pd.notna(tot_dn) else None}
    return df


# --------------------------------------------------------------------------- salidas
def graficos(mensual: pd.DataFrame, out: Path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    etiquetas = [f"{MESES[m - 1]}\n{a}" for a, m in zip(mensual["anio"], mensual["mes"])]
    x = range(len(mensual))
    azul, naranja, tinta, grilla = "#2a6fdb", "#e07b39", "#1f2328", "#d8dde3"

    fig, ax = plt.subplots(figsize=(12, 4.8))
    ax.plot(x, mensual["porcentaje_habitualistas"], color=azul, lw=2.2, marker="o", ms=4)
    for i, v in zip(x, mensual["porcentaje_habitualistas"]):
        ax.annotate(f"{v:.1f}%", (i, v), textcoords="offset points", xytext=(0, 7), ha="center",
                    fontsize=8, color=tinta)
    ax.set_title("Córdoba — % de transferencias que ingresaron a comerciantes habitualistas",
                 loc="left", fontsize=12, color=tinta)
    ax.set_ylabel("% captado por habitualistas")
    ax.set_xticks(list(x), etiquetas, fontsize=8)
    ax.grid(axis="y", color=grilla, lw=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_ylim(0, max(mensual["porcentaje_habitualistas"]) * 1.25)
    fig.text(0.01, 0.01, "Fuente: DNRPA — datos.jus.gob.ar, dataset 'Transferencias de autos'. "
             "Registro seccional en Córdoba; mes según tramite_fecha.", fontsize=7, color="#57606a")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(out / "grafico_pct_habitualistas.png", dpi=150)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(12, 4.8))
    ax.plot(x, mensual["transferencias_totales"], color=azul, lw=2.2, marker="o", ms=4,
            label="Transferencias totales")
    ax.plot(x, mensual["transferencias_habitualistas"], color=naranja, lw=2.2, marker="o", ms=4,
            label="A comerciante habitualista")
    ax.set_title("Córdoba — transferencias totales vs. transferencias a habitualistas",
                 loc="left", fontsize=12, color=tinta)
    ax.set_ylabel("Cantidad de trámites")
    ax.set_xticks(list(x), etiquetas, fontsize=8)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:,.0f}".replace(",", ".")))
    ax.grid(axis="y", color=grilla, lw=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_ylim(0, None)
    ax.legend(frameon=False, loc="upper left")
    fig.text(0.01, 0.01, "Fuente: DNRPA — datos.jus.gob.ar, dataset 'Transferencias de autos'.",
             fontsize=7, color="#57606a")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(out / "grafico_totales_vs_habitualistas.png", dpi=150)
    plt.close(fig)


METODOLOGIA = [
    ("Fuente", "DNRPA / Ministerio de Justicia — datos abiertos, dataset 'Transferencias de autos' "
               "(https://datos.jus.gob.ar/dataset/transferencias-de-autos). Recursos listados vía API CKAN "
               "package_show; ver data/raw/manifest.json (URL, fecha de descarga, SHA-256)."),
    ("Unidad", "Una fila = un trámite de transferencia inscripto en un Registro Seccional (según la metadata "
               "oficial). No hay identificador de dominio, así que no se pueden deduplicar vehículos."),
    ("Filtro geográfico", "registro_seccional_provincia normalizado (mayúsculas, sin tildes) == 'CORDOBA'. "
                          "Es la provincia del Registro Seccional, que según la metadata coincide con el "
                          "domicilio del primer titular o guarda habitual."),
    ("Mes", "Año y mes de tramite_fecha (fecha en que se perfecciona el trámite)."),
    ("Denominador", "Todas las filas de Córdoba del mes, cualquiera sea el tramite_tipo."),
    ("Numerador", f"Filas de Córdoba del mes con tramite_tipo exactamente '{TIPO_HABITUALISTA}'. Otros tipos que "
                  "contengan 'HABITUALISTA' se informan en la hoja de control y NO se suman."),
    ("% captado", "numerador / denominador * 100, por mes. El % acumulado de cada período es "
                  "Σ habitualistas / Σ totales (no el promedio de porcentajes)."),
    ("Qué significa", "Qué proporción de las transferencias inscriptas en Córdoba tuvo como adquirente a un "
                      "comerciante habitualista (vehículo que ingresa formalmente al stock de reventa)."),
    ("Qué NO significa", "No son 'ventas de agencias': no mide lo que las agencias vendieron a consumidores "
                         "finales (esa salida figura como transferencia común desde el habitualista o no se "
                         "registra a su nombre). Tampoco captura compras informales de agencias que no usan "
                         "la figura de habitualista, ni es participación de mercado en unidades distintas."),
    ("Validación externa", "Los totales 2025 aportados por el usuario se usan sólo como comparación; no se "
                           "ajusta ningún dato para que coincidan."),
    ("Actualización", "python analisis_dnrpa/dnrpa_habitualistas_cordoba.py  (descarga sólo archivos nuevos "
                      "y regenera CSV, XLSX, gráficos e informe)."),
]


def escribir_xlsx(path, mensual, resumen, control, inspeccion, validacion, manifest):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    tabla = mensual.copy()
    tabla.insert(0, "Mes", [f"{MESES[m - 1]} {a}" for a, m in zip(tabla["anio"], tabla["mes"])])
    tabla["porcentaje_habitualistas"] = tabla["porcentaje_habitualistas"] / 100

    ctrl_rows = [(k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
                 for k, v in control.items()]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        r = resumen.copy()
        r["pct_acumulado"] = r["pct_acumulado"] / 100
        r["promedio_simple_pct_mensual"] = r["promedio_simple_pct_mensual"] / 100
        r.to_excel(xw, sheet_name="Resumen", index=False)
        tabla.to_excel(xw, sheet_name="Mensual", index=False)
        pd.DataFrame(ctrl_rows, columns=["control", "valor"]).to_excel(xw, sheet_name="Datos - control", index=False)
        fila = len(ctrl_rows) + 3
        validacion.to_excel(xw, sheet_name="Datos - control", index=False, startrow=fila)
        fila += len(validacion) + 3
        pd.DataFrame([{k: (", ".join(v) if isinstance(v, list) else v) for k, v in i.items()}
                      for i in inspeccion]).to_excel(xw, sheet_name="Datos - control", index=False, startrow=fila)
        fila += len(inspeccion) + 3
        pd.DataFrame(manifest).to_excel(xw, sheet_name="Datos - control", index=False, startrow=fila)
        pd.DataFrame(METODOLOGIA, columns=["tema", "detalle"]).to_excel(xw, sheet_name="Metodología", index=False)

        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="2A6FDB")
            for i, col in enumerate(ws.columns, 1):
                ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(12, ancho + 2), 90)
        for ws_name, col in (("Mensual", "F"), ("Resumen", "G"), ("Resumen", "H")):
            for c in xw.book[ws_name][col][1:]:
                c.number_format = "0.00%"
        for ws_name in ("Mensual", "Resumen"):
            for row in xw.book[ws_name].iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, int) and c.number_format == "General":
                        c.number_format = "#,##0"
        for c in xw.book["Metodología"]["B"]:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def fmt_n(n):
    return f"{int(n):,}".replace(",", ".") if pd.notna(n) else "—"


def fmt_p(p):
    return f"{p:.2f}%".replace(".", ",") if pd.notna(p) else "—"


def escribir_informe(path, mensual, resumen, control, inspeccion, validacion, tipos_pais, provincias, manifest,
                     fecha_max):
    L = ["# % captado por habitualistas — Córdoba", "",
         f"Generado: {dt.datetime.now():%Y-%m-%d %H:%M}. Fuente: DNRPA, datos.jus.gob.ar (Transferencias de autos).",
         "", f"**Fecha máxima de tramite_fecha disponible en los archivos procesados:** {fecha_max}", "",
         "## Archivos usados", "", "| archivo | URL | descargado | SHA-256 |", "|---|---|---|---|"]
    L += [f"| {m['archivo']} | {m.get('url') or '(local)'} | {m.get('descargado') or ''} | `{m['sha256'][:16]}…` |"
          for m in manifest]
    i0 = inspeccion[0]
    L += ["", "## Estructura", "",
          f"- Encoding: {sorted({i['encoding'] for i in inspeccion})}; separador: "
          f"{sorted({repr(i['separador']) for i in inspeccion})}",
          f"- Formato tramite_fecha: {sorted({i['formato_tramite_fecha'] for i in inspeccion})}",
          f"- Columnas ({len(i0['columnas'])}): {', '.join(i0['columnas'])}",
          f"- Valores de registro_seccional_provincia que normalizan a CORDOBA: "
          f"{[k for k in provincias if norm(k) == PROVINCIA_OBJETIVO]}",
          "", "### tramite_tipo (todo el país, archivos procesados)", "", "| tramite_tipo | filas |", "|---|---:|"]
    L += [f"| {k} | {fmt_n(v)} |" for k, v in sorted(tipos_pais.items(), key=lambda kv: -kv[1])]
    L += ["", "## Tabla mensual", "", "| Mes | Transferencias totales | A habitualistas | % captado |",
          "|---|---:|---:|---:|"]
    L += [f"| {MESES[r.mes - 1]} {r.anio} | {fmt_n(r.transferencias_totales)} | "
          f"{fmt_n(r.transferencias_habitualistas)} | {fmt_p(r.porcentaje_habitualistas)} |"
          for r in mensual.itertuples()]
    L += ["", "## Resumen", "", "| Período | Meses | Totales | Habitualistas | % acumulado | Promedio simple % mensual |",
          "|---|---:|---:|---:|---:|---:|"]
    L += [f"| {r.periodo} | {r.meses} | {fmt_n(r.transferencias_totales)} | {fmt_n(r.transferencias_habitualistas)} | "
          f"{fmt_p(r.pct_acumulado)} | {fmt_p(r.promedio_simple_pct_mensual)} |" for r in resumen.itertuples()]
    L += ["", "## Controles de calidad", ""]
    L += [f"- **{k}**: {json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v}"
          for k, v in control.items()]
    L += ["", "## Validación externa (2025, sólo sanity check)", "",
          "| Mes | DNRPA datos abiertos | Referencia | Diferencia | Dif. % |", "|---|---:|---:|---:|---:|"]
    L += [f"| {r.mes} | {fmt_n(r.dnrpa_datos_abiertos)} | {fmt_n(r.referencia_externa)} | "
          f"{fmt_n(r.diferencia) if pd.notna(r.diferencia) else '—'} | {fmt_p(r.diferencia_pct)} |"
          for r in validacion.itertuples()]
    L += ["", "## Metodología", ""] + [f"- **{t}.** {d}" for t, d in METODOLOGIA]
    path.write_text("\n".join(L) + "\n", "utf-8")


# --------------------------------------------------------------------------- main
def parse_ym(s):
    a, m = s.split("-")
    return int(a), int(m)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--desde", default="2025-01", type=parse_ym)
    ap.add_argument("--hasta", default="2026-12", type=parse_ym, help="tope; se usa lo publicado hasta ahí")
    ap.add_argument("--solo-local", action="store_true", help="no consultar la API; usar data/raw")
    ap.add_argument("--raw-dir", type=Path, default=RAW)
    ap.add_argument("--out-dir", type=Path, default=OUT)
    ap.add_argument("--eliminar-duplicados", action="store_true",
                    help="descartar filas idénticas (por defecto sólo se informan: sin ID único, "
                         "dos trámites distintos pueden coincidir en todas las columnas)")
    args = ap.parse_args()

    raw, out = args.raw_dir, args.out_dir
    proc = raw.parent / "procesado"
    for d in (raw, proc, out):
        d.mkdir(parents=True, exist_ok=True)

    man_path = raw / "manifest.json"
    manifest_prev = {m["archivo"]: m for m in json.loads(man_path.read_text())} if man_path.exists() else {}

    if not args.solo_local:
        print("Consultando API CKAN…")
        try:
            recursos, meta = listar_recursos(raw)
        except Exception as e:  # red bloqueada, portal caído, etc.
            raise SystemExit(f"No se pudo consultar {CKAN_API}: {e}\n"
                             f"Descargá manualmente los dnrpa-transferencias-autos-AAAAMM.csv desde "
                             f"https://datos.jus.gob.ar/dataset/transferencias-de-autos a {raw} "
                             f"y corré de nuevo con --solo-local.")
        disponibles = sorted((r["anio"], r["mes"] or 0) for r in recursos)
        print(f"  metadata_modified: {meta.get('metadata_modified')}; último recurso: {disponibles[-1]}")
        for r in descargar(recursos, args.desde, args.hasta, raw):
            manifest_prev.setdefault(r["archivo"], {}).update(
                {"archivo": r["archivo"], "url": r["url"], "last_modified_ckan": r["last_modified"]})

    archivos = []
    for p in sorted(raw.iterdir()):
        m = PATRON_ARCHIVO.search(p.name)
        if not m:
            continue
        anio, mes = int(m.group(1)), int(m.group(2)) if m.group(2) else None
        if mes is not None and not (args.desde <= (anio, mes) <= args.hasta):
            continue
        if mes is None and not (args.desde[0] <= anio <= args.hasta[0]):
            continue
        archivos.append(p)
    if not archivos:
        raise SystemExit(f"No hay archivos dnrpa-transferencias-autos-*.csv|zip en {raw}")

    manifest = []
    for p in archivos:
        e = manifest_prev.get(p.name, {"archivo": p.name})
        h = sha256(p)
        if e.get("sha256") != h:
            e["descargado"] = dt.datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds")
        e.update({"sha256": h, "bytes": p.stat().st_size})
        manifest.append(e)
    man_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), "utf-8")

    print("Leyendo archivos…")
    cba, inspeccion, tipos_pais, provincias = procesar_archivos(archivos)
    cba.to_csv(proc / "transferencias_cordoba_filas.csv.gz", index=False, compression="gzip")

    mensual, control, _ = calcular(cba, args.desde, args.hasta, args.eliminar_duplicados)
    fecha_max = max((i["tramite_fecha_max"] for i in inspeccion if i["tramite_fecha_max"]), default=None)
    control["fecha_max_tramite_fecha"] = fecha_max
    ultimo = mensual.iloc[-1]
    if fecha_max and pd.Timestamp(fecha_max).day < 25:
        control["advertencia_ultimo_mes"] = (f"{MESES[int(ultimo.mes) - 1]} {int(ultimo.anio)} podría estar incompleto "
                                             f"(fecha máxima {fecha_max}).")

    resumen = resumen_anual(mensual)
    validacion = validacion_externa(mensual)

    mensual.to_csv(out / "habitualistas_cordoba_mensual.csv", index=False)
    resumen.to_csv(out / "habitualistas_cordoba_resumen.csv", index=False)
    validacion.to_csv(out / "validacion_externa_2025.csv", index=False)
    (out / "control_calidad.json").write_text(json.dumps(
        {"control": control, "inspeccion": inspeccion, "tramite_tipo_pais": tipos_pais,
         "registro_seccional_provincia_valores": provincias}, ensure_ascii=False, indent=2, default=str), "utf-8")
    graficos(mensual, out)
    escribir_xlsx(out / "mercado_habitualistas_cordoba_2025_2026.xlsx", mensual, resumen, control,
                  inspeccion, validacion, manifest)
    escribir_informe(out / "informe.md", mensual, resumen, control, inspeccion, validacion, tipos_pais,
                     provincias, manifest, fecha_max)
    print(f"\nListo. Fecha máxima: {fecha_max}. Salidas en {out}")
    print(mensual.to_string(index=False))
    print(resumen.to_string(index=False))


if __name__ == "__main__":
    sys.exit(main())
